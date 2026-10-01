"""Convert incident markdown report to DOCX and PDF."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from fpdf import FPDF

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "docs" / "INCIDENT-MODULE-ACCESS-OUTAGE-2026-09-18.md"
OUT_DIR = ROOT / "docs"


def strip_md_inline(s: str) -> str:
    s = re.sub(r"\*\*([^*]+)\*\*", r"\1", s)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    return s


def latin1_safe(s: str) -> str:
    repl = {
        "\u2014": "-",
        "\u2013": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2022": "-",
        "\u2192": "->",
        "\u2260": "!=",
        "\u2026": "...",
        "\u00a0": " ",
        "\u2190": "<-",
    }
    for a, b in repl.items():
        s = s.replace(a, b)
    return s.encode("latin-1", "replace").decode("latin-1")


def parse_table_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_separator_row(cells: list[str]) -> bool:
    return all(re.match(r"^:?-+:?$", c.strip() or "-") for c in cells)


def add_runs(paragraph, content: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", content)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def build_docx(lines: list[str], out_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows
        if not table_rows:
            return
        data = [r for r in table_rows if not is_separator_row(r)]
        table_rows = []
        if not data:
            return
        cols = max(len(r) for r in data)
        table = doc.add_table(rows=len(data), cols=cols)
        table.style = "Table Grid"
        for ri, row in enumerate(data):
            for ci in range(cols):
                cell = table.cell(ri, ci)
                cell.text = ""
                p = cell.paragraphs[0]
                val = row[ci].strip() if ci < len(row) else ""
                add_runs(p, val)
                if ri == 0:
                    for run in p.runs:
                        run.bold = True
        doc.add_paragraph()

    for raw in lines:
        line = raw.rstrip()

        if line.startswith("|"):
            table_rows.append(parse_table_row(line))
            continue
        if table_rows:
            flush_table()

        if line == "---" or line.strip() == "":
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- [ ] "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, "[ ] " + line[6:])
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s", "", line))
        elif line.startswith("*") and line.endswith("*") and not line.startswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip("*"))
            run.italic = True
            run.font.size = Pt(10)
        else:
            p = doc.add_paragraph()
            add_runs(p, line)

    if table_rows:
        flush_table()

    doc.save(out_path)


class IncidentPDF(FPDF):
    def footer(self) -> None:
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(100, 100, 100)
        self.cell(0, 10, f"Page {self.page_no()}", align="C")


def write_para(pdf: IncidentPDF, txt: str, size: int = 10, style: str = "", indent: float = 0) -> None:
    pdf.set_font("Helvetica", style, size)
    pdf.set_text_color(30, 30, 30)
    pdf.set_x(pdf.l_margin + indent)
    pdf.multi_cell(0, size * 0.55 + 2.2, latin1_safe(strip_md_inline(txt)))
    pdf.ln(1)


def build_pdf(lines: list[str], out_path: Path) -> None:
    pdf = IncidentPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_margins(18, 18, 18)

    table_buf: list[list[str]] = []

    def flush_pdf_table() -> None:
        nonlocal table_buf
        data = [r for r in table_buf if not is_separator_row(r)]
        table_buf = []
        if not data:
            return
        cols = max(len(r) for r in data)
        usable = pdf.epw
        col_w = usable / cols
        for ri, row in enumerate(data):
            cells = [(row[ci].strip() if ci < len(row) else "") for ci in range(cols)]
            line_counts = []
            for c in cells:
                w = pdf.get_string_width(latin1_safe(strip_md_inline(c)))
                line_counts.append(max(1, int(w / max(col_w - 2, 1)) + 1))
            row_h = max(line_counts) * 4.2 + 2.5
            if pdf.get_y() + row_h > pdf.h - pdf.b_margin:
                pdf.add_page()
            y0 = pdf.get_y()
            x0 = pdf.l_margin
            for ci, c in enumerate(cells):
                x = x0 + ci * col_w
                pdf.rect(x, y0, col_w, row_h)
                pdf.set_xy(x + 1, y0 + 1)
                pdf.set_font("Helvetica", "B" if ri == 0 else "", 8)
                pdf.multi_cell(col_w - 2, 4, latin1_safe(strip_md_inline(c)))
            pdf.set_y(y0 + row_h)
        pdf.ln(3)

    for raw in lines:
        line = raw.rstrip()

        if line.startswith("|"):
            table_buf.append(parse_table_row(line))
            continue
        if table_buf:
            flush_pdf_table()

        if line == "---":
            pdf.ln(2)
            continue
        if line.strip() == "":
            continue

        if line.startswith("# "):
            pdf.ln(4)
            write_para(pdf, line[2:], size=16, style="B")
            pdf.ln(2)
        elif line.startswith("## "):
            pdf.ln(3)
            write_para(pdf, line[3:], size=13, style="B")
            pdf.ln(1)
        elif line.startswith("### "):
            pdf.ln(2)
            write_para(pdf, line[4:], size=11, style="B")
        elif line.startswith("- [ ] "):
            write_para(pdf, "[ ] " + line[6:], size=10, indent=4)
        elif line.startswith("- "):
            write_para(pdf, "- " + line[2:], size=10, indent=4)
        elif re.match(r"^\d+\.\s", line):
            write_para(pdf, line, size=10, indent=4)
        else:
            write_para(pdf, line, size=10)

    if table_buf:
        flush_pdf_table()

    pdf.output(str(out_path))


def main() -> None:
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    docx_path = OUT_DIR / "INCIDENT-MODULE-ACCESS-OUTAGE-2026-09-18.docx"
    pdf_path = OUT_DIR / "INCIDENT-MODULE-ACCESS-OUTAGE-2026-09-18.pdf"
    build_docx(lines, docx_path)
    build_pdf(lines, pdf_path)
    print(f"DOCX: {docx_path}")
    print(f"PDF:  {pdf_path}")


if __name__ == "__main__":
    main()
